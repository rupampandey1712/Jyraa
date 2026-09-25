"""Read a requirements document and turn it into an epic-and-stories plan.

Text extraction is deliberately format-specific rather than a single blob read:
a PDF's page order, a DOCX's headings, and a YAML's structure each carry the
shape of the requirements, and flattening them all the same way loses it.

The planner asks the configured NIM model for structured JSON. When no model is
reachable it falls back to a heading-based split, so uploading a document always
produces something reviewable instead of an error.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
from typing import Any

import yaml

MAX_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 120
MAX_TEXT_CHARS = 120_000

# A model that never answers must not hold the upload request open; past this
# the heading-based plan is returned instead.
# Reasoning models on NIM can take over two minutes to answer a long prompt.
PLAN_TIMEOUT_SECONDS = 210.0

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".csv"}


class DocumentError(ValueError):
    """Raised when a document cannot be read, with a message safe to show a user."""


def extension_of(filename: str) -> str:
    match = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return match.group(1).lower() if match else ""


def extract_text(filename: str, data: bytes) -> str:
    """Pull readable text out of a supported document."""
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")

    suffix = extension_of(filename)
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentError(f"Unsupported file type '{suffix or filename}'. Supported: {supported}.")

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    elif suffix == ".json":
        text = _extract_json(data)
    elif suffix in {".yaml", ".yml"}:
        text = _extract_yaml(data)
    else:
        text = _decode(data)

    text = text.strip()
    if not text:
        raise DocumentError("No readable text was found in the document.")
    return text[:MAX_TEXT_CHARS]


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("The file could not be decoded as text.")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentError("PDF support requires the 'pypdf' package.") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("The PDF could not be opened. It may be corrupt or password protected.") from exc

    pages = []
    for index, page in enumerate(reader.pages):
        if index >= MAX_PDF_PAGES:
            break
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            # One unreadable page should not lose the rest of the document.
            continue

    text = "\n\n".join(part for part in pages if part.strip())
    if not text.strip():
        raise DocumentError(
            "No text layer was found in this PDF. Scanned documents need OCR before they can be read."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise DocumentError("Word support requires the 'python-docx' package.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("The Word file could not be opened. Only .docx is supported, not .doc.") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        # Keep heading levels as markdown so the fallback planner can see structure.
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        match = re.search(r"heading (\d)", style)
        parts.append(f"{'#' * int(match.group(1))} {content}" if match else content)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_json(data: bytes) -> str:
    try:
        parsed = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise DocumentError(f"The JSON file is not valid: {exc.msg} (line {exc.lineno}).") from exc
    return json.dumps(parsed, indent=2, ensure_ascii=False)


def _extract_yaml(data: bytes) -> str:
    try:
        parsed = yaml.safe_load(_decode(data))
    except yaml.YAMLError as exc:
        raise DocumentError(f"The YAML file is not valid: {exc}") from exc
    if parsed is None:
        raise DocumentError("The YAML file is empty.")
    return yaml.safe_dump(parsed, sort_keys=False, allow_unicode=True)


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------

PLAN_SYSTEM = (
    "You are a delivery analyst. Read a requirements document and return only valid JSON. "
    "Do not wrap the response in markdown."
)

PLAN_INSTRUCTIONS = """Convert the requirements below into a delivery plan.

Return JSON shaped exactly like this:
{
  "project_summary": "one or two sentences describing the scope",
  "epics": [
    {
      "summary": "short epic title",
      "description": "what this epic covers and why",
      "stories": [
        {
          "summary": "short story title, written as a deliverable",
          "description": "detail, including acceptance criteria when the document gives them",
          "issue_type": "Story",
          "priority": "Highest|High|Medium|Low|Lowest",
          "estimate_hours": 8,
          "labels": ["optional", "tags"]
        }
      ]
    }
  ]
}

Rules:
- Group work into 1-6 epics. Every story belongs to exactly one epic.
- Use issue_type "Bug" for defect-fixing work, "Task" for chores, otherwise "Story".
- Base everything on the document. Do not invent requirements that are not there.
- estimate_hours is a whole number; omit it when the document gives no basis for one.

Requirements document:
"""


async def plan_from_document(text: str, *, extra_context: str | None = None) -> dict[str, Any]:
    """Ask the model for an epic-and-story breakdown of a requirements document."""
    from app.services.nim_service import generate_response

    prompt = PLAN_INSTRUCTIONS + text
    if extra_context:
        prompt = f"{PLAN_INSTRUCTIONS}Additional instructions from the requester: {extra_context}\n\n{text}"

    try:
        content, model = await asyncio.wait_for(
            generate_response(prompt=prompt, system=PLAN_SYSTEM),
            timeout=PLAN_TIMEOUT_SECONDS,
        )
    except Exception:
        # Timeout, transport error, or a model the account cannot reach.
        content, model = "", None

    plan = _parse_plan(content)
    if plan is not None:
        plan["source"] = "model"
        plan["model"] = model
        return normalize_plan(plan)

    fallback = _plan_from_headings(text)
    fallback["source"] = "fallback"
    fallback["model"] = None
    return normalize_plan(fallback)


def _parse_plan(content: str) -> dict[str, Any] | None:
    """Read the model's JSON, tolerating a stray code fence or surrounding prose."""
    if not content:
        return None
    candidate = content.strip()

    fenced = re.search(r"```(?:json)?\s*(.+?)```", candidate, re.DOTALL)
    if fenced:
        candidate = fenced.group(1).strip()
    else:
        start, end = candidate.find("{"), candidate.rfind("}")
        if start != -1 and end > start:
            candidate = candidate[start:end + 1]

    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        return None
    return parsed if isinstance(parsed, dict) and parsed.get("epics") else None


def _plan_from_headings(text: str) -> dict[str, Any]:
    """Structure a document by its headings when no model is available.

    The shallowest heading level in the document becomes the epic level and
    anything deeper becomes a story, so a document written with ``#``/``##`` and
    one written with ``##``/``###`` both come out the same shape.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    levels = [len(match.group(1)) for match in (re.match(r"^(#{1,6})\s+\S", line) for line in lines) if match]

    epic_level = min(levels) if levels else 1
    title = None
    # A lone top-level heading is the document's title, not an epic. When deeper
    # headings exist, they are the real sections and become the epics instead.
    if levels and levels.count(epic_level) == 1 and any(level > epic_level for level in levels):
        deeper = [level for level in levels if level > epic_level]
        for index, line in enumerate(lines):
            match = re.match(r"^(#{1,6})\s+(.*)$", line)
            if match and len(match.group(1)) == epic_level:
                title = match.group(2).strip()
                lines = lines[:index] + lines[index + 1:]
                break
        epic_level = min(deeper)

    epics: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    preamble: list[str] = []

    def ensure_epic() -> dict[str, Any]:
        nonlocal current
        if current is None:
            current = {"summary": "Requirements", "description": "", "stories": []}
            epics.append(current)
        return current

    for line in lines:
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level, heading_text = len(heading.group(1)), heading.group(2).strip()
            if not heading_text:
                continue
            if level == epic_level:
                current = {"summary": heading_text[:200], "description": "", "stories": []}
                epics.append(current)
            else:
                ensure_epic()["stories"].append({"summary": heading_text[:200], "description": heading_text})
            continue

        # Text before the first section describes the whole document.
        if current is None and not re.match(r"^(?:[-*•]|\d+[.)])\s+", line):
            if len(" ".join(preamble)) < 600:
                preamble.append(line)
            continue

        epic = ensure_epic()
        bullet = re.match(r"^(?:[-*•]|\d+[.)])\s+(.*)$", line)
        if bullet:
            summary = bullet.group(1).strip()
            if summary:
                epic["stories"].append({"summary": summary[:200], "description": summary})
        elif len(epic["description"]) < 1000:
            epic["description"] = f"{epic['description']} {line}".strip()

    # An epic the document never broke down still needs one story to be actionable.
    for epic in epics:
        if not epic["stories"]:
            epic["stories"].append({
                "summary": f"Deliver {epic['summary']}"[:200],
                "description": epic["description"] or epic["summary"],
            })

    parts = [part for part in (title, " ".join(preamble).strip()) if part]
    parts.append("Structured from the document's headings; no model was available.")

    return {"project_summary": " ".join(parts), "epics": epics[:6]}


_PRIORITIES = {"highest", "high", "medium", "low", "lowest"}
_ISSUE_TYPES = {"story", "task", "bug"}


def normalize_plan(plan: dict[str, Any]) -> dict[str, Any]:
    """Clamp a plan to values the issue API will accept."""
    epics: list[dict[str, Any]] = []

    for raw_epic in (plan.get("epics") or [])[:10]:
        if not isinstance(raw_epic, dict):
            continue
        summary = str(raw_epic.get("summary") or "").strip()
        if not summary:
            continue

        stories: list[dict[str, Any]] = []
        for raw_story in (raw_epic.get("stories") or [])[:50]:
            if not isinstance(raw_story, dict):
                continue
            story_summary = str(raw_story.get("summary") or "").strip()
            if not story_summary:
                continue

            issue_type = str(raw_story.get("issue_type") or "Story").strip().lower()
            priority = str(raw_story.get("priority") or "Medium").strip().lower()
            estimate = raw_story.get("estimate_hours")
            try:
                estimate = float(estimate) if estimate is not None else None
            except (TypeError, ValueError):
                estimate = None

            stories.append({
                "summary": story_summary[:200],
                "description": str(raw_story.get("description") or "").strip()[:4000],
                "issue_type": issue_type.title() if issue_type in _ISSUE_TYPES else "Story",
                "priority": priority.title() if priority in _PRIORITIES else "Medium",
                "estimate_hours": estimate if estimate and estimate > 0 else None,
                "labels": [str(label).strip()[:50] for label in (raw_story.get("labels") or [])[:10] if str(label).strip()],
            })

        epics.append({
            "summary": summary[:200],
            "description": str(raw_epic.get("description") or "").strip()[:4000],
            "stories": stories,
        })

    return {
        "project_summary": str(plan.get("project_summary") or "").strip()[:2000],
        "epics": epics,
        "source": plan.get("source", "model"),
        "model": plan.get("model"),
        "story_count": sum(len(epic["stories"]) for epic in epics),
    }
